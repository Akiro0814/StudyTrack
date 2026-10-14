#!/usr/bin/env python3
"""Synthetic, deterministic fixtures produced independently with openpyxl/stdlib."""
import copy, hashlib, json, pathlib, zipfile
from datetime import datetime
from openpyxl import Workbook
from openpyxl.utils.datetime import CALENDAR_MAC_1904

root = pathlib.Path(__file__).resolve().parents[1]
fixtures = root / 'StudyTrackTests' / 'Fixtures'
fixtures.mkdir(parents=True,exist_ok=True)
headers = ['学期','科目','考试名称','时间','分数','满分','分类','班排','年排','区排','反思']
book = Workbook(); sheet = book.active; sheet.title = 'Grades'; sheet.append(headers)
sheet.append(['秋季','数学','=文字', '2026-10-12',89.5,100,'月考',None,1,None,'中文\n第二行'])
sheet['C2'].data_type='s'
book.save(fixtures / 'independent.xlsx')
book = Workbook(); sheet=book.active
sheet.append(['反思','分数','区排','时间','科目','年排','分类','学期','满分','考试名称','班排'])
sheet.append(['第一行\n第二行',85,None,61,'数学',1,'月考','秋季',100,'月考',None]); sheet['D2'].number_format='yyyy-mm-dd'
sheet.append([None,60,None,60,'数学',None,'期末','秋季',100,'非法日期',None]); sheet['D3'].number_format='yyyy-mm-dd'
sheet.append([None,'=1+2',None,'2026-10-12','数学',None,'月考','秋季',100,'公式',None])
sheet.append([None,80,None,'2026-02-30','数学',None,'月考','秋季',100,'错误日期',None])
sheet.append([None,80.001,None,'2026-10-12','数学',None,'月考','秋季',100,'超精度',None])
sheet.append([None,80,None,datetime(2026,10,12,12,30),'数学',None,'周考','秋季',100,'待确认',None]); sheet['D6'].number_format='yyyy-mm-dd hh:mm'
book.save(fixtures / 'edge-cases.xlsx')
book = Workbook(); book.epoch=CALENDAR_MAC_1904; sheet=book.active; sheet.append(headers)
sheet.append(['春季','数学','期中',0,120,150,'期中',None,1,None,None]); sheet['D2'].number_format='yyyy-mm-dd'
book.save(fixtures / 'date1904.xlsx')
book = Workbook(); sheet=book.active; sheet.append(['科目','时间','分数','科目']); sheet.append(['数学','2026-10-12',85,'数学'])
book.save(fixtures / 'duplicate-headers.xlsx')

def uid(n): return f'00000000-0000-4000-8000-{n:012d}'
stamp='2026-10-12T00:00:00.123Z'
def common(n,name): return dict(id=uid(n),name=name,createdAt=stamp,updatedAt=stamp)
users=[common(1,'学生甲'),common(2,'学生乙')]
terms=[dict(common(10+i,'秋季' if i%2==0 else '春季'),userID=uid(1+i//2)) for i in range(4)]
subjects=[dict(common(20,'数学'),userID=uid(1),semesterID=uid(10),fullScore='100.00'),
          dict(common(21,'物理'),userID=uid(1),semesterID=uid(10),fullScore='150.00'),
          dict(common(22,'空科目'),userID=uid(1),semesterID=uid(11),fullScore='100.00'),
          dict(common(23,'数学'),userID=uid(2),semesterID=uid(12),fullScore='150.00')]
exams=[]
for n,subject,score in [(30,20,'80.00'),(31,21,'120.00')]:
    exams.append(dict(common(n,'期中'),userID=uid(1),semesterID=uid(10),subjectID=uid(subject),date='2026-10-12',score=score,type='midterm',classRank=None,gradeRank=1 if n==31 else None,districtRank=None,reflection='中文\n反思' if n==30 else None,attachmentID=None))
data=dict(users=users,semesters=terms,subjects=subjects,examRecords=exams,assets=[],legacyIDMap=[])
preferences=dict(theme='glass',accent='blue',multiUserEnabled=True,currentUserID=uid(1),currentSemesterID=uid(10),wallpaperAssetID=None,nameNormalizationVersion=1,platformExtensions={'android':{'unknown':{'nested':[1,True,'保留']},'sourceTimeZone':'Asia/Shanghai'}})
def save_backup(name,data=data,preferences=preferences,extra=None,version=1):
    files={'data.json':json.dumps(data,ensure_ascii=False).encode(),'preferences.json':json.dumps(preferences,ensure_ascii=False).encode()}
    if extra: files.update(extra)
    manifest=dict(format='studytrack-backup',formatVersion=version,backupID=uid(99),sourceDatasetID=uid(100),sourcePlatform='android',sourceAppVersion='test-fixture',exportedAt=stamp,recordCounts={key:len(value) for key,value in data.items()},files=[dict(path=key,byteCount=len(value),sha256=hashlib.sha256(value).hexdigest()) for key,value in files.items()])
    files['manifest.json']=json.dumps(manifest).encode()
    with zipfile.ZipFile(fixtures/name,'w',zipfile.ZIP_DEFLATED) as z:
        for path,value in files.items(): z.writestr(path,value)
save_backup('golden.studytrackbackup')
save_backup('unknown-version.studytrackbackup',version=2)
save_backup('unsafe-path.studytrackbackup',extra={'../escape':b'x'})
broken=copy.deepcopy(data); broken['examRecords'][0]['subjectID']=uid(999)
save_backup('broken-reference.studytrackbackup',data=broken)
(fixtures/'expected.json').write_text(json.dumps({'users':2,'semesters':4,'subjects':4,'examRecords':2,'scope':{'userID':uid(1),'semesterID':uid(10)},'average':100,'minimum':80,'maximum':120,'excellenceRate':0},ensure_ascii=False,indent=2)+'\n')
print('Generated independent XLSX/backup fixtures in',fixtures)
